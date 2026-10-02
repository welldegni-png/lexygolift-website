from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


SOURCE = Path(
    r"F:\business and work\sale detail\长兴圣力\LEXYGO\CODEX拓展业务\叉车+堆高车\堆高车宣传册制作\WEB"
)
OUTPUT = Path(__file__).resolve().parents[1] / "source-data.json"


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = " ".join(value.replace("\u3000", " ").split())
        return value or None
    return value


def extract_workbook(path: Path):
    workbook = load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for worksheet in workbook.worksheets:
        cells = []
        for row in worksheet.iter_rows():
            row_values = []
            for cell in row:
                value = clean(cell.value)
                if value is not None:
                    row_values.append({"cell": cell.coordinate, "value": value})
            if row_values:
                cells.append(row_values)
        sheets.append(
            {
                "title": worksheet.title,
                "cells": cells,
                "merged": [str(value) for value in worksheet.merged_cells.ranges],
            }
        )
    return sheets


def extract_document(path: Path):
    document = Document(path)
    paragraphs = [clean(p.text) for p in document.paragraphs]
    tables = []
    for table in document.tables:
        rows = []
        for row in table.rows:
            rows.append([clean(cell.text) for cell in row.cells])
        tables.append(rows)
    return {
        "paragraphs": [p for p in paragraphs if p],
        "tables": tables,
    }


def main():
    payload = {"source": str(SOURCE), "workbooks": [], "documents": []}
    for path in sorted(SOURCE.rglob("*.xlsx")):
        payload["workbooks"].append(
            {
                "path": str(path.relative_to(SOURCE)),
                "sheets": extract_workbook(path),
            }
        )
    for path in sorted(SOURCE.rglob("*.docx")):
        payload["documents"].append(
            {
                "path": str(path.relative_to(SOURCE)),
                **extract_document(path),
            }
        )
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Extracted {len(payload['workbooks'])} workbooks and {len(payload['documents'])} documents")
    print(OUTPUT)


if __name__ == "__main__":
    main()
